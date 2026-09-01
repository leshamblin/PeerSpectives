import os
from docx import Document
from docx.oxml.ns import qn

def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    p._p = p._element = None

def delete_table(table):
    t = table._element
    t.getparent().remove(t)
    t._tbl = t._element = None

def get_team_name(text):
    # Expecting "Team [Name] – Team Peer Review Stats"
    # Returns "Team [Name]"
    if "–" in text:
        return text.split("–")[0].strip()
    if "-" in text:
        return text.split("-")[0].strip()
    return text.strip()

def split_teams(source_path):
    if not os.path.exists(source_path):
        print(f"Error: File not found: {source_path}")
        return

    # 1. Identify Teams and their content indices
    # We'll scan the document to find "Team ..." headers.
    # We assume the structure is Header -> Table.
    
    doc = Document(source_path)
    teams = []
    
    current_team = None
    
    # We need to iterate over body elements to get order
    # doc.paragraphs and doc.tables are separate lists.
    # Let's map paragraphs and tables to their index in the body?
    # Or simpler: Just iterate paragraphs, find the header, then find the *next* table.
    
    # Let's try a robust approach:
    # Find all paragraphs that look like headers.
    # For each header, find the table that follows it immediately (ignoring empty paragraphs).
    
    header_paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("Team ") and "Peer Review Stats" in p.text:
            header_paragraphs.append((i, p.text))
            
    print(f"Found {len(header_paragraphs)} teams.")
    
    for i, (p_index, header_text) in enumerate(header_paragraphs):
        team_name = get_team_name(header_text)
        safe_team_name = team_name.replace(" ", "_")
        output_filename = f"{safe_team_name}_Stats.docx"
        
        print(f"Processing {team_name} -> {output_filename}")
        
        # Reload fresh doc for each team to delete others
        team_doc = Document(source_path)
        
        # We need to find the specific paragraph object in the new doc that corresponds to our header
        # and the specific table that follows it.
        
        # Since we are reloading, indices should be stable for paragraphs/tables lists *if* we don't modify.
        # But we want to modify.
        
        # Strategy:
        # 1. Identify the index of the header paragraph in `doc.paragraphs`.
        # 2. Identify the index of the target table in `doc.tables`.
        #    The target table is the first table *after* the header paragraph in the document flow.
        
        # Let's map document flow to find the table index.
        # This is tricky in python-docx.
        
        # Alternative Strategy:
        # Iterate over `team_doc.element.body` children.
        # Keep the child if it matches the header or the table we want.
        # Delete otherwise.
        
        # To match: use text for paragraph.
        # For table: this is harder.
        # Assumption: The table immediately follows the header.
        
        body_elements = team_doc.element.body
        
        elements_to_keep = []
        
        # State machine scan
        found_header = False
        kept_table = False
        
        # We need to collect elements to remove, then remove them.
        elements_to_remove = []
        
        for element in body_elements:
            # Check if it's a paragraph
            if element.tag.endswith('p'):
                text = ""
                # Extract text from xml directly to be safe or wrap in Paragraph
                # Simple text extraction from xml
                for node in element.iter():
                    if node.tag.endswith('t'):
                        if node.text:
                            text += node.text
                
                if text.strip() == header_text.strip():
                    found_header = True
                    continue # Keep this
                
            # Check if it's a table
            if element.tag.endswith('tbl'):
                if found_header and not kept_table:
                    kept_table = True
                    continue # Keep this (it's the table after the header)
            
            # If we haven't decided to keep it, mark for removal
            elements_to_remove.append(element)
            
            # Reset found_header if we've found both (so we don't keep subsequent tables)
            if found_header and kept_table:
                found_header = False 
                # We don't reset kept_table because we only want one set per file?
                # Actually, for this loop, we only want ONE team.
                # So once we have kept the table, we are done finding things to keep.
        
        # Remove elements
        for el in elements_to_remove:
            el.getparent().remove(el)
            
        team_doc.save(output_filename)
        print(f"Saved {output_filename}")

if __name__ == "__main__":
    split_teams("Team_Stats_Tables_With_Summary.docx")
